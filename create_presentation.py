from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

ROOT = Path(__file__).parent
OUT = ROOT / "Dynamic_Indoor_Navigation_Presentation.pptx"
ASSETS = ROOT / "presentation_assets"
ASSETS.mkdir(exist_ok=True)

NAVY = RGBColor(17, 48, 78); BLUE = RGBColor(35, 111, 176); TEAL = RGBColor(16, 143, 142)
ORANGE = RGBColor(231, 128, 47); RED = RGBColor(199, 71, 71); GREEN = RGBColor(54, 150, 103)
INK = RGBColor(35, 46, 56); MUTED = RGBColor(92, 106, 120); WHITE = RGBColor(255, 255, 255); LIGHT = RGBColor(243, 247, 250)
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5); blank = prs.slide_layouts[6]

def bg(s, c=WHITE): s.background.fill.solid(); s.background.fill.fore_color.rgb = c
def box(s,x,y,w,h,c,round=True):
    z=s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if round else MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)); z.fill.solid(); z.fill.fore_color.rgb=c; z.line.color.rgb=c; return z
def tx(s,t,x,y,w,h,size=16,c=INK,b=False,align=PP_ALIGN.LEFT):
    z=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); f=z.text_frame; f.clear(); f.word_wrap=True; f.margin_left=f.margin_right=Inches(.05); p=f.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=t; r.font.name='Aptos'; r.font.size=Pt(size); r.font.bold=b; r.font.color.rgb=c; return z
def heading(s,t,sub=''):
    tx(s,t,.65,.34,12,.5,28,NAVY,True); box(s,.66,.94,1.05,.045,TEAL,False)
    if sub: tx(s,sub,.68,1.03,11.8,.25,11,MUTED)
def foot(s,n):
    box(s,0,7.23,13.333,.27,NAVY,False); tx(s,'DYNAMIC INDOOR NAVIGATION USING PDR AND ROUTE OPTIMIZATION',.55,7.255,10,.15,7.5,WHITE,True); tx(s,str(n),12.25,7.25,.4,.15,8,WHITE,True,PP_ALIGN.RIGHT)
def card(s,x,y,w,h,hdr,body,accent=TEAL):
    box(s,x,y,w,h,LIGHT); box(s,x,y,.1,h,accent); tx(s,hdr,x+.3,y+.2,w-.5,.3,15,NAVY,True); tx(s,body,x+.3,y+.67,w-.55,h-.8,11.5,INK)
def bullet_lines(s,x,y,items,accent=TEAL):
    for i,(h,b) in enumerate(items):
        yy=y+i*.83; box(s,x,yy,.36,.36,accent); tx(s,str(i+1),x,yy+.08,.36,.15,9,WHITE,True,PP_ALIGN.CENTER); tx(s,h,x+.55,yy-.02,3.6,.2,12,NAVY,True); tx(s,b,x+.55,yy+.25,3.8,.28,10.5,MUTED)

def demo_route():
    import warnings; warnings.filterwarnings('ignore')
    from app import app
    c=app.test_client(); png=c.post('/get_path',json={'start':'G28','end':'G04'}); out=ASSETS/'route_demo.png'
    if png.status_code==200: out.write_bytes(png.data)
    js=c.post('/get_path',json={'start':'G28','end':'G04','format':'json'}).get_json() or {}
    return out,js

image, route = demo_route()

# 1 Title
s=prs.slides.add_slide(blank); bg(s,NAVY); box(s,8.9,0,4.43,7.5,RGBColor(19,59,97),False); box(s,9.7,0,.1,7.5,TEAL,False)
tx(s,'DYNAMIC\nINDOOR NAVIGATION',.8,1.15,7.3,1.3,31,WHITE,True); tx(s,'Using PDR and Route Optimization',.84,2.8,6.8,.4,19,RGBColor(184,227,225),True)
tx(s,'A mobile-first system that maps a user’s movement indoors and selects efficient, congestion-aware routes.',.85,3.7,6.9,.7,15,WHITE)
tx(s,'PROJECT PRESENTATION',.86,5.94,2.3,.2,10,RGBColor(184,227,225),True); tx(s,'R. Akshya  •  Samriddhi Shaw  •  Meghana Gopinath Tanuja',.86,6.28,7.3,.25,12,WHITE)
for y,label,c in [(1.2,'START',GREEN),(2.6,'PDR',ORANGE),(4.0,'ROUTE',TEAL),(5.4,'DESTINATION',BLUE)]: box(s,10.25,y,2,.55,c); tx(s,label,10.3,y+.16,1.9,.2,11,WHITE,True,PP_ALIGN.CENTER)

# 2 Problem
s=prs.slides.add_slide(blank); bg(s); heading(s,'Problem Statement','Indoor navigation needs to work where GPS is unreliable or unavailable.')
tx(s,'Finding a room, exit, or service inside a multi-floor building can be slow, confusing, and unsafe during congestion or emergencies.',.72,1.45,11.9,.58,18,INK)
card(s,.75,2.35,3.85,2.65,'GPS DOES NOT SOLVE INDOORS','Satellite signals are weakened by walls and floors, so a phone cannot reliably locate or guide a user inside a building.',BLUE)
card(s,4.75,2.35,3.85,2.65,'STATIC MAPS ARE NOT ENOUGH','Printed maps and fixed directions do not adapt to a user’s movement, floor changes, or the nearest emergency exit.',ORANGE)
card(s,8.75,2.35,3.85,2.65,'CROWDS CHANGE THE BEST ROUTE','The shortest physical path can become inconvenient or unsafe when corridors are congested.',RED)
tx(s,'Goal: provide accurate, real-time and adaptive indoor guidance using a smartphone and a digital floor-map graph.',.82,5.6,11.5,.35,16,TEAL,True); foot(s,2)

# 3 Proposed solution
s=prs.slides.add_slide(blank); bg(s); heading(s,'Proposed Solution','A dynamic navigation pipeline that fuses pedestrian motion with mapped building connectivity.')
for i,(num,h,b,c) in enumerate([('1','SELECT','Start room / destination',BLUE),('2','LOCATE','PDR step + map matching',TEAL),('3','OPTIMIZE','A* shortest or crowd-aware route',ORANGE),('4','GUIDE','Floor-wise visual route',GREEN)]):
    x=.7+i*3.15; box(s,x,2.08,2.5,2.3,LIGHT); box(s,x+.2,2.26,.55,.55,c); tx(s,num,x+.2,2.4,.55,.16,15,WHITE,True,PP_ALIGN.CENTER); tx(s,h,x+.22,3.08,2,.25,15,NAVY,True); tx(s,b,x+.22,3.55,2,.5,12,INK)
tx(s,'Key output',.77,5.3,1.2,.2,13,NAVY,True); tx(s,'A navigable path that can span floors via mapped staircases and avoid simulated high-density areas.',1.95,5.3,9.7,.25,14,INK); foot(s,3)

# 4 Architecture
s=prs.slides.add_slide(blank); bg(s); heading(s,'Three-Tier Architecture','Separation of mobile interaction, application services and spatial intelligence makes the system extensible.')
tiers=[('PRESENTATION TIER','Flutter mobile app',['Start / destination selection','PDR controls: heading + step length','Route map, crowd simulation & status'],BLUE),('APPLICATION TIER','Flask REST API',['/pdr_step and /map_match','/get_path and /simulate_crowd','Room matching + response formatting'],TEAL),('DATA & INTELLIGENCE TIER','Spatial graph engine',['GeoJSON floor plans (3 levels)','NetworkX corridor/stair graph','A* + crowd-weighted edge costs'],ORANGE)]
for i,(a,b,items,c) in enumerate(tiers):
    y=1.5+i*1.7; box(s,.75,y,11.85,1.3,LIGHT); box(s,.75,y,.15,1.3,c); tx(s,a,1.15,y+.18,2.7,.2,11,c,True); tx(s,b,1.15,y+.52,2.7,.25,16,NAVY,True)
    for j,item in enumerate(items): box(s,4.15+j*2.75,y+.34,2.45,.6,WHITE); tx(s,item,4.25+j*2.75,y+.49,2.25,.2,9.5,INK,False,PP_ALIGN.CENTER)
tx(s,'Data flows downward for computation; route and status updates flow back to the user.',.8,6.55,11.6,.2,12,MUTED,False,PP_ALIGN.CENTER); foot(s,4)

# 5 Working
s=prs.slides.add_slide(blank); bg(s); heading(s,'How Navigation Works','Phone-based motion estimates are constrained to valid indoor walking space.')
box(s,.75,1.55,5.65,4.9,LIGHT); tx(s,'PDR + MAP MATCHING',1.02,1.85,4.8,.3,17,NAVY,True); bullet_lines(s,1.05,2.55,[('Inputs','Current position, heading and step length'),('PDR step','Compute the next raw latitude/longitude estimate'),('Map match','Snap the estimate to the nearest corridor geometry')])
box(s,6.9,1.55,5.65,4.9,LIGHT); tx(s,'ROUTE OPTIMIZATION',7.18,1.85,4.8,.3,17,NAVY,True); bullet_lines(s,7.2,2.55,[('Build graph','Corridor boundaries become weighted graph edges'),('Link floors','Matched staircase nodes connect levels'),('Choose route','A* minimizes distance or crowd-weighted cost')],ORANGE); foot(s,5)

# 6 Novelty
s=prs.slides.add_slide(blank); bg(s); heading(s,'Novelty of Our Work','A practical integration of positioning, spatial routing and dynamic occupancy awareness.')
items=[('01','Map-constrained PDR','Estimated pedestrian steps are snapped to nearby mapped corridors, reducing impossible off-path movement.',TEAL),('02','Multi-floor graph routing','Three floor plans are combined through matched staircase connections for end-to-end navigation.',BLUE),('03','Crowd-aware path cost','Simulated or schedule-derived occupancy is flood-filled across corridors and penalizes crowded edges.',RED),('04','Low-infrastructure design','The demonstrator uses mobile inputs, GeoJSON maps and software crowd simulation—without dedicated beacons.',ORANGE)]
for i,(n,h,b,c) in enumerate(items):
    x=.75+(i%2)*6.05; y=1.65+(i//2)*2.4; box(s,x,y,5.7,1.8,LIGHT); box(s,x+.25,y+.27,.7,.7,c); tx(s,n,x+.25,y+.49,.7,.16,14,WHITE,True,PP_ALIGN.CENTER); tx(s,h,x+1.18,y+.28,4.1,.25,16,NAVY,True); tx(s,b,x+1.18,y+.75,4.08,.6,11.5,INK)
foot(s,6)

# 7 Demo
s=prs.slides.add_slide(blank); bg(s); heading(s,'Demo: Dynamic Route Generation','Example using the implemented web service and the project’s mapped ground floor.')
if image.exists(): s.shapes.add_picture(str(image), Inches(.75), Inches(1.45), width=Inches(7.1), height=Inches(5.48))
else: box(s,.75,1.45,7.1,5.48,LIGHT)
box(s,8.2,1.45,4.35,5.48,LIGHT); tx(s,'DEMO FLOW',8.55,1.78,3.5,.25,14,TEAL,True)
bullet_lines(s,8.55,2.3,[('Enter locations','G28 → G04'),('Find shortest route',str(route.get('distance_m','—'))+' m computed by A*'),('Try PDR','Heading + step length → matched corridor point'),('Simulate crowd','Route around hotspots or apply timetable occupancy')],BLUE)
tx(s,'Live output: route map, distance, crowd exposure and hotspot count.',8.55,6.1,3.4,.4,11.5,INK,True); foot(s,7)

# 8 Future
s=prs.slides.add_slide(blank); bg(s); heading(s,'Future Scope','The current prototype provides a foundation for richer positioning, accessibility and deployment features.')
future=[('Sensor fusion','Fuse IMU data with Wi-Fi, BLE beacons, UWB or visual markers to improve positioning accuracy.'),('Continuous calibration','Estimate stride length and heading drift per user; correct position using landmarks and turns.'),('Live operations','Connect crowd data to cameras, turnstiles, Wi-Fi analytics or IoT sensors for real-time congestion.'),('Accessible routing','Offer wheelchair-friendly paths, elevator preferences, evacuation routes and multilingual voice guidance.')]
for i,(h,b) in enumerate(future):
    x=.75+(i%2)*6.1; y=1.65+(i//2)*2.25; box(s,x,y,5.55,1.72,WHITE); box(s,x,y,.11,1.72,[TEAL,BLUE,ORANGE,GREEN][i]); tx(s,h,x+.35,y+.28,4.8,.25,16,NAVY,True); tx(s,b,x+.35,y+.78,4.72,.55,11.5,INK)
tx(s,'Long-term vision: an adaptive digital twin for safe, efficient movement inside campuses, hospitals, malls and public buildings.',.8,6.35,11.8,.3,13,TEAL,True,PP_ALIGN.CENTER); foot(s,8)

# 9 Closing
s=prs.slides.add_slide(blank); bg(s,NAVY); box(s,8.9,0,4.43,7.5,RGBColor(19,59,97),False); tx(s,'Thank You',.75,1.35,7.4,.6,36,WHITE,True); tx(s,'Dynamic Indoor Navigation Using PDR and Route Optimization',.8,2.35,7.8,.4,18,RGBColor(184,227,225),True); tx(s,'Questions?',.8,3.35,2.3,.35,19,WHITE,True)
for y,h,c in [(1.5,'PDR',TEAL),(2.75,'MAP MATCH',BLUE),(4.0,'A* ROUTING',ORANGE),(5.25,'CROWD AWARE',RED)]: box(s,9.55,y,2.2,.56,c); tx(s,h,9.63,y+.16,2.03,.18,11,WHITE,True,PP_ALIGN.CENTER)

prs.save(OUT)
print(OUT)
